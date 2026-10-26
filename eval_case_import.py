"""Spreadsheet parsing, templates and annotated error files for evaluation imports."""
import base64
import csv
import io
import json
import re
import zipfile
from pathlib import PurePosixPath

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_UPLOAD = 12 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 120
MAX_EXPANDED = 60 * 1024 * 1024
TEMPLATE_URL = 'https://nwd4iy9rd2s.feishu.cn/wiki/WUBjwVB08i7WXAkEvGycO6Mrn0e?sheet=f71768'
TEMPLATE_HEADERS = ['用例 ID（必填，唯一）', 'Stage（必填）', 'Story（必填）', 'prompt_EN（必填）',
                    'prompt_CN（必填）', 'Skill', 'Factor 1', 'Factor 取值 1', 'Factor 2',
                    'Factor 取值 2', 'Factor 3', 'Factor 取值 3', '所属项目（必填）',
                    '数据归属（默认 Eval-白盒）', '道具', '图片', '可见性', '可见用户组',
                    'Skill 类别', '标签', '图片类型']


def normalized(value):
    value = re.sub(r'[（(](?:必填|默认)[^）)]*[）)]', '', str(value or ''))
    return re.sub(r'[\s_*＊]', '', value).casefold()


def cell_text(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, bool):
        return '是' if value else '否'
    if isinstance(value, int) and re.fullmatch(r'0{2,}', cell.number_format or ''):
        return str(value).zfill(len(cell.number_format))
    return str(value)


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED:
            raise ValueError('文件解压后过大，请拆分文件后重试')


def read_book(data, filename):
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix == '.xlsx':
        check_archive(data)
        return load_workbook(io.BytesIO(data), data_only=False)
    if suffix == '.csv':
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('gb18030')
        book = Workbook()
        sheet = book.active
        sheet.title = '评测用例'
        for number, values in enumerate(csv.reader(io.StringIO(text)), 1):
            if number > MAX_ROWS + 20 or len(values) > MAX_COLUMNS:
                raise ValueError('单次最多导入 5000 条用例、120 列')
            for column, value in enumerate(values, 1):
                cell = sheet.cell(number, column, value)
                cell.data_type = 's'
        return book
    raise ValueError('请选择 .xlsx 或 .csv 文件')


def parse_upload(data, filename):
    if not data:
        raise ValueError('文件为空，请重新选择')
    if len(data) > MAX_UPLOAD:
        raise ValueError('文件大小不能超过 12 MB')
    book = read_book(data, filename)
    candidates = []
    for sheet in book.worksheets:
        if sheet.title in {'填写说明', '校验说明'}:
            continue
        if sheet.max_column > MAX_COLUMNS or sheet.max_row > MAX_ROWS + 20:
            raise ValueError('单次最多导入 5000 条用例、120 列')
        for number, row in enumerate(sheet.iter_rows(max_row=min(15, sheet.max_row)), 1):
            names = {normalized(cell.value) for cell in row}
            if names.intersection({'用例id', 'caseid'}) and names.intersection({'prompten', 'prompt'}):
                candidates.append((sheet, number))
                break
    if len(candidates) != 1:
        raise ValueError('未找到唯一的评测用例表头，请使用指定模板，保留用例 ID、Stage、Story、prompt_EN、prompt_CN 表头')
    sheet, header_row = candidates[0]
    headers = [cell_text(cell).strip() for cell in sheet[header_row]]
    # Existing error reports can be fixed and uploaded again.
    error_columns = {index for index, value in enumerate(headers) if normalized(value) in {'错误原因', '导入错误'}}
    columns = [index for index in range(len(headers)) if index not in error_columns]
    last = max((index for index in columns if headers[index]), default=-1)
    columns = [index for index in columns if index <= last]
    headers = [headers[index] for index in columns]
    rows, issues = [], []
    for area in sheet.merged_cells.ranges:
        if area.max_row >= header_row:
            issues.append({'row': header_row, 'message': f'「{area}」存在合并单元格，请取消合并并逐行填写'})
    notes = '\n'.join(cell_text(cell) for row in sheet.iter_rows(max_row=header_row - 1) for cell in row)
    sample_match = re.search(r'第\s*(\d+)\s*行[^\n；;。]*示例[^\n；;。]*不导入', notes)
    sample_row = int(sample_match.group(1)) if sample_match else None
    for number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
        if number == sample_row:
            continue
        values = [cell_text(cells[index]) for index in columns]
        if not any(value.strip() for value in values):
            continue
        for index in columns:
            if cells[index].data_type == 'f':
                issues.append({'row': number, 'message': f'{cell_text(sheet.cell(header_row, index + 1))} 不支持公式，请粘贴为值'})
            elif cells[index].data_type == 'e' or isinstance(cells[index].value, bool) or cells[index].is_date:
                issues.append({'row': number, 'message': f'{cell_text(sheet.cell(header_row, index + 1))} 不是有效文本，请检查单元格内容'})
            if normalized(sheet.cell(header_row, index + 1).value) in {'用例id', 'caseid'}:
                value = cells[index].value
                if isinstance(value, (int, float)) and not isinstance(value, bool) and (not float(value).is_integer() or abs(value) >= 10**15):
                    issues.append({'row': number, 'field': 'id', 'message': '用例 ID 可能已被 Excel 改写，请将该列设置为文本并填写完整 ID'})
        rows.append({'row': number, 'values': values})
    if len(rows) > MAX_ROWS:
        raise ValueError('单次最多导入 5000 条用例')
    embedded = {}
    image_columns = {index + 1 for index, value in enumerate([cell_text(cell) for cell in sheet[header_row]])
                     if normalized(value) == '图片'}
    for image in sorted(sheet._images, key=lambda item: (getattr(getattr(item.anchor, '_from', None), 'row', 0), getattr(getattr(item.anchor, '_from', None), 'rowOff', 0), getattr(getattr(item.anchor, '_from', None), 'colOff', 0))):
        anchor = getattr(image.anchor, '_from', None)
        if anchor is None or anchor.col + 1 not in image_columns or anchor.row + 1 <= header_row:
            issues.append({'row': anchor.row + 1 if anchor else header_row, 'message': '请将图片放在「图片」列对应数据行内'})
            continue
        if anchor.row + 1 > MAX_ROWS + 20:
            raise ValueError('图片所在行超出单次导入范围，请将图片放在对应的用例数据行内')
        try:
            raw = image._data()
            item = image_item(raw, '嵌入图片.' + image.format)
            embedded.setdefault(str(anchor.row + 1), []).append(item)
        except ValueError as error:
            issues.append({'row': anchor.row + 1, 'message': str(error)})
    populated = {row['row'] for row in rows}
    for number in embedded:
        number = int(number)
        if number not in populated and number != sample_row:
            rows.append({'row': number, 'values': [''] * len(headers)})
    rows.sort(key=lambda row: row['row'])
    if len(rows) > MAX_ROWS:
        raise ValueError('单次最多导入 5000 条用例')
    return {'filename': filename, 'sheet': sheet.title, 'header_row': header_row, 'headers': headers, 'rows': rows, 'issues': issues, 'embedded': embedded}


def image_item(raw, name):
    from PIL import Image
    if len(raw) > 1024 * 1024:
        raise ValueError(f'{name} 超过单张图片 1 MB 限制')
    try:
        with Image.open(io.BytesIO(raw)) as image:
            image.verify()
            if image.format not in {'PNG', 'JPEG', 'WEBP', 'GIF'}:
                raise ValueError('不支持的图片格式')
            actual_mime = Image.MIME[image.format]
    except Exception as error:
        raise ValueError(f'{name} 图片损坏或格式不支持') from error
    return {'name': name, 'src': f'data:{actual_mime};base64,' + base64.b64encode(raw).decode()}


def error_workbook(data, filename, errors, sheet_name='', header_row=1):
    try:
        book = read_book(data, filename)
        sheet = book[sheet_name] if sheet_name in book.sheetnames else book.active
        if sheet.max_row > MAX_ROWS + 20 or sheet.max_column > MAX_COLUMNS:
            raise ValueError('文件超出导入范围')
    except Exception:
        book = Workbook()
        sheet = book.active
        sheet.title = '导入错误'
        sheet.append(['文件名'])
        sheet.append([filename])
        header_row = 1
        errors = [{'row': 2, 'message': item.get('message', '')} for item in errors]
    header_row = max(1, min(int(header_row), sheet.max_row))
    existing = [cell.column for cell in sheet[header_row] if normalized(cell.value) in {'错误原因', '导入错误'}]
    column = existing[0] if existing else sheet.max_column + 1
    for row in sheet.iter_rows(min_col=column, max_col=column):
        row[0].value = None
    cell = sheet.cell(header_row, column, '错误原因')
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='C62828')
    sheet.column_dimensions[cell.column_letter].width = 70
    grouped = {}
    if '校验说明' in book.sheetnames:
        del book['校验说明']
    guide = book.create_sheet('校验说明')
    guide.append(['本批次未导入任何数据。错误原因为空仅表示该行未发现错误，不表示已导入。'])
    guide.column_dimensions['A'].width = 100
    for item in errors:
        row = max(header_row, min(int(item.get('row') or header_row), MAX_ROWS + 20))
        grouped.setdefault(row, []).append(str(item.get('message', '校验失败')))
    for row, messages in grouped.items():
        if row == header_row:
            guide.append(['；'.join(dict.fromkeys(messages))])
            continue
        cell = sheet.cell(row, column, '\n'.join(dict.fromkeys(messages)))
        cell.data_type = 's'
        cell.font = Font(color='C62828')
        cell.fill = PatternFill('solid', fgColor='FFF1F0')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for worksheet in book:
        for row in worksheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def register_case_import(app):
    def template():
        book = Workbook()
        sheet = book.active
        sheet.title = '导入模板'
        sheet.append(TEMPLATE_HEADERS)
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='168C9B')
            sheet.column_dimensions[cell.column_letter].width = 28 if cell.column not in (4, 5) else 55
        guide = book.create_sheet('填写说明')
        for row in [
            ['字段', '填写与校验规则'],
            ['模板来源', TEMPLATE_URL],
            ['用例 ID', '必填，仅字母、数字、下划线、短横线；文件内必须唯一。与平台已有 ID 重复时，按抽屉所选方式更新或跳过（默认跳过）'],
            ['重复处理', '更新：按表内字段覆盖，非必填项留空会清空，未提供的列保留原值；发布/启用状态、创建信息保留，记录实际修改。跳过：该行不写入，也不创建关联数据'],
            ['Stage / Story', '必填，仅在所属项目下按完整名称或 ID 匹配；不存在则新增为已发布、启用；其他项目同名不影响新建；已有未发布或停用要素报错'],
            ['prompt_EN / prompt_CN', '均必填，不支持公式，不自动翻译'],
            ['Skill', '非必填，中文（英文），每行一个；仅在所属项目下匹配。新增必须在 Skill 类别列逐行填写现有类别，如动作类/辨别类；自动新建为已发布、启用'],
            ['Factor N / Factor 取值 N', '同编号成对填写，Factor 不得重复；仅在所属项目下匹配。不存在时按「一级 > 二级 > 三级」自动新建（三级可空）；未知取值追加到该 Factor。多值每行一个；含逗号的单个新值用双引号包裹，已有完整值优先匹配'],
            ['所属项目', '必填，预训练评测 / 后训练评测；新建关联要素沿用该项目'],
            ['数据归属', '留空默认 Eval-白盒，支持 Train-训练 / Eval-白盒 / Test-黑盒；其他值自动加入自定义选项'],
            ['图片', '直接在 Excel「图片」列对应用例行插入图片，无需单独上传图片文件；每例最多 3 张、每张 ≤1MB；导入后默认初始状态，可在编辑中逐张调整类型'],
            ['图片类型', '可选，按同格图片由上到下、从左到右的顺序逐行填写初始状态/目标状态/关键物体位置；为空默认初始状态'],
            ['标签', '可选，填写已有完整标签路径或 ID，多项每行一个'],
            ['字符限制', 'ID ≤100；Stage/Story/Skill 名称及 Factor 每级 ≤200；Factor 单个取值 ≤500；prompt_EN/CN 各 ≤10000；道具 ≤2000；数据归属 ≤100'],
            ['可见性 / 可见用户组', '可见性默认公开；受限时填写已有用户组名称或 ID，多选每行一个；公开时用户组留空'],
            ['导入结果', '新用例未发布、未启用；整批校验通过才写入，用例及关联要素、场景组合一起保存；失败不写入，返回错误原因列'],
        ]:
            guide.append(row)
        guide.column_dimensions['A'].width = 30
        guide.column_dimensions['B'].width = 110
        for row in guide:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        stream = io.BytesIO()
        book.save(stream)
        stream.seek(0)
        return send_file(stream, as_attachment=True, download_name='评测用例导入模板.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def uploaded():
        file = request.files.get('file')
        if not file or not file.filename:
            raise ValueError('请选择导入文件')
        data = file.stream.read(MAX_UPLOAD + 1)
        if len(data) > MAX_UPLOAD:
            raise ValueError('文件大小不能超过 12 MB')
        return data, file.filename

    def parse():
        result = None
        try:
            data, filename = uploaded()
            result = parse_upload(data, filename)
            return jsonify(result)
        except Exception as error:
            return jsonify(message=str(error) if isinstance(error, ValueError) else '文件无法读取，请检查表格格式或内嵌图片是否损坏',
                           sheet=result.get('sheet', '') if result else '',
                           header_row=result.get('header_row', 1) if result else 1), 400

    def report():
        try:
            data, filename = uploaded()
            errors = json.loads(request.form.get('errors', '[]'))
            if not isinstance(errors, list) or len(errors) > 50000:
                raise ValueError('错误信息格式无效')
            stream = error_workbook(data, filename, errors, request.form.get('sheet', ''), request.form.get('header_row', 1))
            return send_file(stream, as_attachment=True, download_name=PurePosixPath(filename).stem + '_错误明细.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        except Exception as error:
            return jsonify(message=str(error) if isinstance(error, ValueError) else '错误文件生成失败，请重试'), 400

    app.add_url_rule('/model/eval/test-cases/import/template', 'eval_case_import_template', template)
    app.add_url_rule('/model/eval/test-cases/import/parse', 'eval_case_import_parse', parse, methods=['POST'])
    app.add_url_rule('/model/eval/test-cases/import/errors', 'eval_case_import_errors', report, methods=['POST'])
    from eval_import_jobs import register_import_jobs
    register_import_jobs(app)
