import io
import json
import unittest

from flask import Flask
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from eval_case_import import TEMPLATE_HEADERS, error_workbook, image_item, parse_upload, register_case_import


def workbook_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def fixture():
    book = Workbook()
    sheet = book.active
    sheet.title = '导入模板'
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(['IMPORT_01', '厨房餐厅', '餐桌整理', 'Pick.', '拿起。', '', '', '', '', '', '', '', '预训练评测'])
    return book


class EvalCaseImportTests(unittest.TestCase):
    def setUp(self):
        app = Flask(__name__)
        register_case_import(app)
        self.client = app.test_client()

    def test_live_headers_and_hidden_columns(self):
        book = fixture()
        book.active.column_dimensions['R'].hidden = True
        book.active['R2'] = 'QA'
        parsed = parse_upload(workbook_bytes(book), 'template.xlsx')
        self.assertEqual(parsed['header_row'], 1)
        self.assertEqual(parsed['headers'], TEMPLATE_HEADERS)
        self.assertEqual(parsed['rows'][0]['row'], 2)
        self.assertEqual(parsed['rows'][0]['values'][17], 'QA')
        self.assertEqual(parsed['issues'], [])

    def test_formulas_and_zero_padded_ids(self):
        book = fixture()
        book.active['A2'] = 12
        book.active['A2'].number_format = '0000'
        book.active['D2'] = '=1+1'
        parsed = parse_upload(workbook_bytes(book), 'x.xlsx')
        self.assertEqual(parsed['rows'][0]['values'][0], '0012')
        self.assertEqual(parsed['issues'][0]['row'], 2)
        self.assertIn('公式', parsed['issues'][0]['message'])

    def test_offset_headers_explicit_sample_and_empty_rows(self):
        book = fixture()
        book.active.insert_rows(1, 2)
        book.active['A1'] = '第4行是示例，不导入'
        book.active['A6'] = 'IMPORT_02'
        book.active['D6'] = 'Other prompt'
        parsed = parse_upload(workbook_bytes(book), 'x.xlsx')
        self.assertEqual(parsed['header_row'], 3)
        self.assertEqual([row['row'] for row in parsed['rows']], [6])

    def test_csv_keeps_text_and_row_numbers(self):
        parsed = parse_upload('用例 ID,Stage,Story,prompt_EN,prompt_CN,所属项目\n0001,厨房,任务,Pick.,拿起。,预训练评测\n'.encode('utf-8-sig'), 'x.csv')
        self.assertEqual(parsed['rows'][0]['values'][0], '0001')
        self.assertEqual(parsed['rows'][0]['row'], 2)
        with self.assertRaises(ValueError):
            parse_upload(b'invalid', 'x.xls')
        with self.assertRaises(ValueError):
            parse_upload(b'', 'x.xlsx')

    def test_report_keeps_all_sheets_and_all_failed_rows(self):
        book = fixture()
        book.create_sheet('填写说明')['A1'] = 'Original guide'
        book.active.append(['IMPORT_02'])
        data = workbook_bytes(book)
        errors = [{'row': 2, 'message': 'ID 重复'}, {'row': 2, 'message': 'Story 必填'}, {'row': 3, 'message': 'ID 重复'}, {'row': 1, 'message': '缺少字段'}]
        report = error_workbook(data, 'x.xlsx', errors, '导入模板', 1)
        annotated = load_workbook(report)
        self.assertEqual(annotated['导入模板'].cell(1, len(TEMPLATE_HEADERS)+1).value, '错误原因')
        self.assertEqual(annotated['导入模板'].cell(2, len(TEMPLATE_HEADERS)+1).value, 'ID 重复\nStory 必填')
        self.assertEqual(annotated['导入模板'].cell(3, len(TEMPLATE_HEADERS)+1).value, 'ID 重复')
        self.assertEqual(annotated['填写说明']['A1'].value, 'Original guide')
        fixed = parse_upload(workbook_bytes(annotated), 'corrected.xlsx')
        self.assertNotIn('错误原因', fixed['headers'])
        # File-level errors are annotations, never synthetic evaluation rows.
        self.assertEqual(len(fixed['rows']), 2)

    def test_embedded_images(self):
        png = io.BytesIO()
        Image.new('RGB', (10, 10), 'white').save(png, format='PNG')
        book = fixture()
        book.active.add_image(SheetImage(io.BytesIO(png.getvalue())), 'P2')
        parsed = parse_upload(workbook_bytes(book), 'x.xlsx')
        self.assertEqual(len(parsed['embedded']['2']), 1)
        self.assertTrue(parsed['embedded']['2'][0]['src'].startswith('data:image/png;base64,'))
        with self.assertRaisesRegex(ValueError, '损坏'):
            image_item(b'invalid image', 'bad.png')

    def test_routes_download_parse_and_annotate(self):
        response = self.client.get('/model/eval/test-cases/import/template')
        self.assertEqual(response.status_code, 200)
        template = load_workbook(io.BytesIO(response.data))
        self.assertEqual([cell.value for cell in template['导入模板'][1]], TEMPLATE_HEADERS)
        self.assertEqual(template['导入模板'].max_row, 1, 'Download has no sample that could be accidentally imported')
        data = workbook_bytes(fixture())
        response = self.client.post('/model/eval/test-cases/import/parse', data={'file': (io.BytesIO(data), 'x.xlsx')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['rows'][0]['row'], 2)
        response = self.client.post('/model/eval/test-cases/import/errors', data={'file': (io.BytesIO(data), 'x.xlsx'), 'sheet': '导入模板', 'header_row': '1', 'errors': json.dumps([{'row': 2, 'message': 'ID 已存在'}])})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(response.data)).active.cell(2, len(TEMPLATE_HEADERS)+1).value, 'ID 已存在')
        response = self.client.post('/model/eval/test-cases/import/parse', data={'file': (io.BytesIO(b'broken'), 'x.xlsx')})
        self.assertEqual(response.status_code, 400)
        response = self.client.post('/model/eval/test-cases/import/errors', data={'file': (io.BytesIO(b'broken'), 'x.xlsx'), 'errors': json.dumps([{'row': 1, 'message': '文件损坏'}])})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(response.data)).active['B2'].value, '文件损坏')


if __name__ == '__main__':
    unittest.main()
