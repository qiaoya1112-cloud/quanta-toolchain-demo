import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook

from eval_catalog import seed_data
from eval_case_import import TEMPLATE_HEADERS, register_case_import
from eval_import_jobs import connect, run_catalog

BASE = '/model/eval/test-cases/import'


def upload_rows(rows):
    book = Workbook()
    book.active.title = '导入模板'
    book.active.append(TEMPLATE_HEADERS)
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def valid_row(case_id='IMPORT_JOB_1', project='预训练评测'):
    return [case_id, '导入场域', '导入任务', 'Pick up the cup.', '拿起杯子。',
            '导入能力（Import skill）', '导入因素 > 几何属性 > 形状', '新值', '', '', '', '',
            project, '', '', '', '公开', '', '动作类', '', '']


class EvalImportJobsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.app = Flask(__name__)
        self.app.secret_key = 'test-only'
        self.app.config.update(TESTING=True, EVAL_IMPORT_SYNC=True,
                               EVAL_IMPORT_DB=str(Path(self.temp.name)/'imports.sqlite3'),
                               EVAL_IMPORT_CONTEXT_FACTORY=lambda: {'userGroups': [], 'tags': []})
        register_case_import(self.app)
        self.a, self.b = self.app.test_client(), self.app.test_client()
        for client, user in [(self.a, 'A'), (self.b, 'B')]:
            with client.session_transaction() as session:
                session['tenant_id'] = 'tenant'
                session['user_id'] = user
                session['user_name'] = 'User ' + user
        result = self.a.post(BASE+'/state', json={'document': seed_data(), 'revision': 0})
        self.assertEqual(result.status_code, 200)

    def submit(self, rows, key='request-1', mode='skip', client=None):
        return (client or self.a).post(BASE+'/jobs', data={
            'file': (io.BytesIO(upload_rows(rows)), '评测.xlsx'),
            'request_key': key, 'duplicate_mode': mode})

    def state(self):
        return self.a.get(BASE+'/state').json

    def test_project_uniqueness_all_entities_and_value_reuse(self):
        response = self.submit([valid_row(), valid_row('IMPORT_JOB_2', '后训练评测')])
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.json['status'], 'success', response.json)
        data = self.state()['document']
        for kind, field, value in [('stages','name','导入场域'),('stories','name','导入任务'),
                                   ('skills','name_zh','导入能力'),('factors','category','导入因素')]:
            matches = [x for x in data[kind] if x.get(field)==value]
            self.assertEqual(len(matches), 2, kind)
            self.assertEqual({x['project'] for x in matches}, {'预训练评测','后训练评测'})
            self.assertEqual(len({x['id'] for x in matches}), 2)
            self.assertTrue(all(x['enabled'] and x['publish_status']=='已发布' for x in matches))
            self.assertTrue(all(x['created_by']=='User A' for x in matches))
        factor = next(x for x in data['factors'] if x.get('category')=='导入因素' and x['project']=='预训练评测')
        added = valid_row('IMPORT_JOB_3'); added[7]='新值\n另一个值'
        result = self.submit([added], key='next').json
        self.assertEqual(result['status'], 'success', result)
        self.assertEqual(result['summary']['related']['factors'], 0)
        self.assertEqual(result['summary']['related']['factor_values'], 1)
        after = next(x for x in self.state()['document']['factors'] if x['id']==factor['id'])
        self.assertEqual([x['value'] for x in after['values']], ['新值','另一个值'])
        self.assertEqual(after['default_value'], '')
        self.assertTrue(after['update_history'])

    def test_failure_atomic_and_download_isolation(self):
        before = self.state()
        bad = valid_row('BAD'); bad[4]='';bad[16]='受限'
        response = self.submit([valid_row(), bad])
        task = response.json
        self.assertEqual(task['status'], 'validation_failed', task)
        self.assertEqual(self.state(), before)
        file = self.a.get(task['download_url'])
        self.assertEqual(file.status_code, 200)
        book = load_workbook(io.BytesIO(file.data))
        self.assertEqual(book.active.cell(1,len(TEMPLATE_HEADERS)+1).value,'错误原因')
        self.assertIn('prompt_CN',book.active.cell(3,len(TEMPLATE_HEADERS)+1).value)
        self.assertIn('可见用户组',book.active.cell(3,len(TEMPLATE_HEADERS)+1).value)
        self.assertEqual(self.b.get(BASE+'/jobs').json['total'],0)
        self.assertEqual(self.b.get(task['download_url']).status_code,404)
        self.assertEqual(self.b.get(BASE+'/jobs/'+task['id']).status_code,404)
        # Client query parameters cannot change ownership.
        self.assertEqual(self.b.get(BASE+'/jobs?user_id=A').json['total'],0)
        self.assertEqual(self.a.get(BASE+'/jobs').json['items'][0]['status'],'validation_failed')

    def test_idempotency_updates_skips_and_audit(self):
        response=self.submit([valid_row()]); task=response.json
        self.assertEqual(task['status'],'success',task)
        same=self.submit([valid_row()]).json
        self.assertEqual(same['id'],task['id'])
        self.assertEqual(self.a.get(BASE+'/jobs').json['total'],1)
        replacement=valid_row();replacement[4]='更新后的提示词。'
        conflict=self.submit([replacement])
        self.assertEqual(conflict.status_code,409)
        skip=self.submit([replacement],key='skip').json
        self.assertEqual(skip['summary']['skipped'],1)
        changed=self.submit([replacement],key='update',mode='update').json
        self.assertEqual(changed['summary']['updated'],1)
        case=next(x for x in self.state()['document']['test-cases'] if x['id']=='IMPORT_JOB_1')
        self.assertEqual(case['prompt_cn'],'更新后的提示词。')
        self.assertEqual(case['publish_status'],'未发布');self.assertFalse(case['enabled'])
        unchanged=self.submit([replacement],key='unchanged',mode='update').json
        self.assertEqual(unchanged['summary']['unchanged'],1)

    def test_inactive_in_project_does_not_fall_back_to_other_project(self):
        self.submit([valid_row(),valid_row('SECOND','后训练评测')])
        state=self.state();doc=state['document']
        stage=next(x for x in doc['stages'] if x['name']=='导入场域' and x['project']=='预训练评测')
        stage['enabled']=False
        self.assertEqual(self.a.post(BASE+'/state',json={'document':doc,'revision':state['revision']}).status_code,200)
        result=self.submit([valid_row('THIRD')],key='inactive').json
        self.assertEqual(result['status'],'validation_failed')
        book=load_workbook(io.BytesIO(self.a.get(result['download_url']).data))
        self.assertIn('已停用',book.active.cell(2,len(TEMPLATE_HEADERS)+1).value)

    def test_concurrent_catalog_change_aborts_whole_import(self):
        original=run_catalog
        def race(app,payload):
            result=original(app,payload)
            if payload['action']=='plan':
                with connect(app) as db:
                    db.execute('UPDATE eval_catalog_state SET revision=revision+1')
            return result
        before=self.state()['document']
        with patch('eval_import_jobs.run_catalog', side_effect=race):
            task=self.submit([valid_row()]).json
        self.assertEqual(task['status'],'import_failed')
        self.assertEqual(self.state()['document'],before)

    def test_persisted_history_after_recreating_app(self):
        task=self.submit([valid_row()]).json
        other=Flask('restart');other.secret_key='test-only';other.config.update(self.app.config)
        register_case_import(other);client=other.test_client()
        with client.session_transaction() as session:
            session['tenant_id']='tenant';session['user_id']='A'
        self.assertEqual(client.get(BASE+'/jobs').json['items'][0]['id'],task['id'])
        self.assertEqual(client.get(task['download_url']).status_code,200)

    def reserve(self, key='staged', filename='评测.xlsx'):
        response = self.a.post(BASE+'/jobs', json={'filename':filename,'request_key':key,'duplicate_mode':'skip'})
        self.assertEqual(response.status_code,202)
        return response.json

    def test_upload_lifecycle_retry_isolation_and_timeout(self):
        task=self.reserve()
        self.assertEqual(task['status'],'uploading')
        self.assertIsNone(task['download_url'])
        endpoint=BASE+'/jobs/'+task['id']+'/upload'
        self.assertEqual(self.b.post(endpoint,json={}).status_code,404)
        self.assertEqual(self.a.get(BASE+'/jobs/'+task['id']+'/file').status_code,404)
        failed=self.a.post(endpoint,json={'reason':'empty'}).json
        self.assertEqual(failed['status'],'upload_failed')
        self.assertIn('文件为空',failed['message'])
        retry=self.reserve()
        self.assertEqual(retry['id'],task['id'])
        self.assertEqual(retry['status'],'uploading')
        self.assertEqual(self.a.get(BASE+'/jobs').json['total'],1)
        uploaded=self.a.post(endpoint,data={'file':(io.BytesIO(upload_rows([valid_row()])), '评测.xlsx')}).json
        self.assertEqual(uploaded['status'],'success',uploaded)
        self.assertEqual(self.a.post(endpoint,json={}).json['status'],'success')
        self.assertEqual(self.a.get(uploaded['download_url']).status_code,200)
        interrupted=self.reserve(key='interrupted')
        with connect(self.app) as db:
            db.execute('UPDATE eval_import_jobs SET lease_until=0 WHERE id=?',(interrupted['id'],))
        history=self.a.get(BASE+'/jobs').json['items']
        self.assertEqual(next(row for row in history if row['id']==interrupted['id'])['status'],'upload_failed')

    def test_upload_validating_importing_and_success_transitions(self):
        from contextlib import contextmanager
        from eval_import_jobs import worker
        task=self.reserve()
        observed=[]
        @contextmanager
        def observe(app):
            with connect(app) as db:
                yield db
                observed.append(db.execute('SELECT status FROM eval_import_jobs WHERE id=?',(task['id'],)).fetchone()['status'])
        with patch('eval_import_jobs.start'):
            response=self.a.post(BASE+'/jobs/'+task['id']+'/upload',data={'file':(io.BytesIO(upload_rows([valid_row()])), '评测.xlsx')})
        self.assertEqual(response.json['status'],'validating')
        with patch('eval_import_jobs.connect',side_effect=observe):
            worker(self.app,task['id'])
        self.assertIn('validating',observed)
        self.assertIn('importing',observed)
        self.assertEqual(observed[-1],'success')

    def test_uploaded_invalid_file_validation_failure_report(self):
        before=self.state()
        task=self.reserve(filename='broken.xlsx')
        result=self.a.post(BASE+'/jobs/'+task['id']+'/upload',data={'file':(io.BytesIO(b'invalid workbook'), 'broken.xlsx')}).json
        self.assertEqual(result['status'],'validation_failed')
        self.assertTrue(result['has_report'])
        self.assertEqual(self.a.get(result['download_url']).status_code,200)
        self.assertEqual(self.state(),before)


if __name__ == '__main__':
    unittest.main()
