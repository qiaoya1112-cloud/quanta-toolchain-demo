"""XLSX transport for browser-local instruction package creation."""
from io import BytesIO
from flask import jsonify, request
from openpyxl import load_workbook


def register_batch_packages(app):
    @app.post('/data/instruction-packages/batch-parse')
    def batch_package_parse():
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.xlsx'):
            return jsonify(error='仅支持上传 .xlsx 文件'), 400
        data = upload.read(20 * 1024 * 1024 + 1)
        if len(data) > 20 * 1024 * 1024:
            return jsonify(error='文件大小不能超过 20 MB'), 400
        book = None
        try:
            book = load_workbook(BytesIO(data), read_only=True, data_only=False)
            sheet = book.worksheets[0]
            values = sheet.iter_rows(values_only=True)
            headers = [str(v or '').strip().replace(' ', '') for v in next(values, [])]
            if len(set(headers)) != len(headers):
                raise ValueError('表头存在重复列')
            if '指令包序号' not in headers or '指令ID' not in headers:
                raise ValueError('缺少必填列：指令包序号、指令 ID')
            rows = []
            for line, row in enumerate(values, 2):
                if line > 50001:
                    raise ValueError('最多支持 50,000 行，请拆分文件')
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                record = {'line': line}
                for key in ('指令包序号', '指令ID', '指令版本', '指令包名称'):
                    value = row[headers.index(key)] if key in headers and headers.index(key) < len(row) else ''
                    record[key] = str(value if value is not None else '').strip()
                rows.append(record)
            if not rows:
                raise ValueError('表格没有数据，请填写后重新上传')
            return jsonify(rows=rows)
        except ValueError as error:
            return jsonify(error=str(error)), 400
        except Exception:
            return jsonify(error='无法读取 XLSX，请检查文件是否损坏、加密或格式不正确'), 400
        finally:
            if book:
                book.close()
